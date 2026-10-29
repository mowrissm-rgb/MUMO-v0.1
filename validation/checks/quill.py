"""Quill check: every value in a generated report equals its run evidence.

Protocol: validation/PROTOCOL.md, section "Quill: report fidelity".
"""
import io
import json
import math

UNAVAILABLE_WORDS = ('unavailable', 'not calculated')
BAD_MISSING_TEXT = {'', 'none', 'nan', '0', '0.0', 'null'}


def _tables(docx_bytes):
    from docx import Document
    doc = Document(io.BytesIO(docx_bytes))
    tables = [[[cell.text for cell in row.cells] for row in table.rows]
              for table in doc.tables]
    paragraphs = [p.text for p in doc.paragraphs]
    return tables, paragraphs


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _same(cell, value):
    """Cell text equals the evidence value (numbers compared numerically)."""
    if _is_number(value):
        try:
            return float(cell) == float(value)
        except ValueError:
            return False
    return cell == str(value)


def _missing_ok(cell):
    text = cell.strip().lower()
    return text not in BAD_MISSING_TEXT and any(w in text for w in UNAVAILABLE_WORDS)


def _kv_rows(tables):
    """All two-column rows from every table, keyed by label (first wins)."""
    rows = {}
    for table in tables:
        for row in table:
            if len(row) == 2:
                rows.setdefault(row[0], row[1])
    return rows


def _check_druglikeness(docx_bytes, evidence_dl, memory_dl):
    """Compare the drug-likeness table with the evidence copy. Returns (matched, total, misses)."""
    tables, _ = _tables(docx_bytes)
    rows = _kv_rows(tables)
    matched, misses = 0, []
    fields = {k: v for k, v in evidence_dl.items() if not k.startswith('_')}
    for key, value in fields.items():
        cell = rows.get(key)
        if cell is None:
            misses.append({'field': key, 'problem': 'row missing from report'})
        elif value is None:
            if _missing_ok(cell):
                matched += 1
            else:
                misses.append({'field': key, 'problem': f'missing value rendered as {cell!r}'})
        elif _same(cell, value):
            matched += 1
        else:
            misses.append({'field': key, 'problem': f'report {cell!r} != evidence {value!r}'})
    # evidence JSON must also equal the in-memory values it was built from
    for key, value in memory_dl.items():
        if key.startswith('_'):
            continue
        ev = evidence_dl.get(key)
        mem_missing = value is None or (isinstance(value, float) and not math.isfinite(value))
        same = ev is None if mem_missing else (_same(str(ev), value) if _is_number(value) else ev == value)
        if not same:
            misses.append({'field': key, 'problem': f'evidence {ev!r} != input {value!r}'})
    return matched, len(fields), misses


def _check_metabolism(docx_bytes, evidence_pred):
    tables, paragraphs = _tables(docx_bytes)
    mets = evidence_pred.get('metabolites') or []
    total = len(mets)
    misses = []
    header = ['Rank', 'Metabolite (SMILES)', 'Transformation', 'Phase', 'Score', 'Route']
    table = next((t for t in tables if t and t[0] == header), None)
    count_line = f"{len(mets)} metabolites shown"
    if not any(count_line in p for p in paragraphs):
        misses.append({'field': 'count line', 'problem': f'{count_line!r} not in report'})
    if not mets:
        if table is not None:
            misses.append({'field': 'table', 'problem': 'table present for zero metabolites'})
        return 0, 0, misses
    if table is None:
        return 0, total, [{'field': 'table', 'problem': 'metabolite table missing'}]
    body = table[1:]
    if len(body) != total:
        misses.append({'field': 'rows', 'problem': f'{len(body)} rows != {total} metabolites'})
    matched = 0
    for rank, (met, row) in enumerate(zip(mets, body), 1):
        expected = [rank, met['smiles'], met['name'], met['phase'], met['score'],
                    ' -> '.join(met['pathway'])]
        wrong = [h for h, cell, exp in zip(header, row, expected) if not _same(cell, exp)]
        if wrong:
            misses.append({'field': f'metabolite {rank}', 'problem': f'columns differ: {wrong}'})
        else:
            matched += 1
    return matched, total, misses


def _evidence(result):
    from verification import evidence_json
    return json.loads(evidence_json(result))


def run(manifest, ctx):
    import report_writer
    from agents.admet import druglikeness
    from agents.metabolism import predict_metabolites

    items = []
    totals = {'matched': 0, 'fields': 0}
    fixture = manifest['admet_ai_fixture']

    def record(name, matched, total, misses, error=None):
        totals['matched'] += matched
        totals['fields'] += total
        if error:
            status = 'ERROR'
        else:
            status = 'PASS' if not misses and matched == total else 'FAIL'
        items.append({'item': name, 'status': status,
                      'details': {'fields_matched': matched, 'fields_total': total,
                                  'misses': misses, 'error': error}})

    for drug in manifest['drugs']:
        name, smiles = drug['name'], drug['smiles']
        try:
            dl = druglikeness(smiles)
            pred = predict_metabolites(smiles)
            admet_result = {'lig_label': name, 'lig_smiles': smiles,
                            'druglikeness': dl, 'admet_ml': dict(fixture)}
            met_result = {'lig_label': name, 'lig_smiles': smiles, 'prediction': pred}
            admet_docx = report_writer.build_admet_docx(admet_result)
            met_docx = report_writer.build_metabolism_docx(met_result)
            ev_admet = _evidence(admet_result)
            ev_met = _evidence(met_result)
            m1, t1, miss1 = _check_druglikeness(admet_docx, ev_admet['druglikeness'], dl)
            m2, t2, miss2 = _check_metabolism(met_docx, ev_met['prediction'])
            record(name, m1 + m2, t1 + t2, miss1 + miss2)
        except Exception as exc:  # errors stay in the denominator
            record(name, 0, 0, [], error=f'{type(exc).__name__}: {exc}')

    inj = manifest['injection']
    dl = {k: (float('nan') if v == 'NaN' else v) for k, v in inj['druglikeness'].items()}
    try:
        result = {'lig_label': inj['name'], 'lig_smiles': inj['smiles'], 'druglikeness': dl}
        docx = report_writer.build_admet_docx(result)
        m, t, misses = _check_druglikeness(docx, _evidence(result)['druglikeness'], dl)
        record(f"injection:{inj['name']}", m, t, misses)
    except Exception as exc:
        record(f"injection:{inj['name']}", 0, 0, [], error=f'{type(exc).__name__}: {exc}')

    n_pass = sum(1 for item in items if item['status'] == 'PASS')
    criteria = [{
        'id': 'Q1',
        'description': '100% of fields match on 100% of items (20 drugs + 1 injection)',
        'passed': n_pass == len(items) and totals['matched'] == totals['fields'],
        'value': f"{n_pass}/{len(items)} items pass; {totals['matched']}/{totals['fields']} fields match",
    }]
    extras = {'fields_matched': totals['matched'], 'fields_total': totals['fields'],
              'admet_ai_fixture_only': True}
    return items, criteria, extras
