"""Five-panel target figures on one A4 portrait report page.

FreeSASA runs in a child process, never in the shared Streamlit environment.
The stored cleaned receptor is the common input for geometry, area and ribbon.
"""
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import tempfile


def solvent_exposure(pdb_text, timeout=60):
    python = os.environ.get('MUMO_VALIDATION_PYTHON', '/opt/validationenv/bin/python')
    if not Path(python).is_file():
        return {'status': 'unavailable', '_error': 'FreeSASA interpreter unavailable'}
    with tempfile.TemporaryDirectory(prefix='mumo-sasa-') as tmp:
        source = Path(tmp) / 'receptor.pdb'
        source.write_text(pdb_text, encoding='utf-8', newline='')
        try:
            run = subprocess.run([python, str(Path(__file__).with_name('sasa_runner.py')), str(source)],
                                 capture_output=True, text=True, timeout=timeout)
            if run.returncode:
                raise RuntimeError((run.stderr or 'FreeSASA failed')[-600:])
            result = json.loads(run.stdout)
            if result.get('status') != 'calculated':
                raise ValueError(result.get('_error', 'Invalid FreeSASA result'))
            if result['input_sha256'] != hashlib.sha256(pdb_text.encode('utf-8')).hexdigest():
                raise ValueError('Solvent exposure input provenance mismatch')
            return result
        except (OSError, subprocess.SubprocessError, ValueError, KeyError, RuntimeError) as exc:
            return {'status': 'unavailable', '_error': str(exc)}


def compose_panels(panels):
    from PIL import Image, ImageDraw, ImageFont, ImageChops
    canvas = Image.new('RGB', (2200, 2700), 'white')
    draw = ImageDraw.Draw(canvas)
    try:
        font = ImageFont.truetype('DejaVuSans.ttf', 32)
    except OSError:
        font = ImageFont.load_default(size=32)
    layout = [
        ('protein', 'A  Protein structure', (0, 0, 700, 530)),
        ('exposure', 'B  Solvent exposure', (750, 0, 700, 530)),
        ('count', 'C  Residue count', (1500, 0, 700, 530)),
        ('rama', 'D  Ramachandran plot', (150, 560, 1900, 1500)),
        ('polar', 'E  Polar and apolar surface', (0, 2100, 2200, 600)),
    ]
    for key, label, (x, y, w, h) in layout:
        draw.text((x + 12, y + 4), label, font=font, fill='black')
        png = panels.get(key)
        if not png:
            draw.text((x + 12, y + 100), 'Not calculated', font=font, fill='#667085')
            continue
        image = Image.open(io.BytesIO(png)).convert('RGB')
        mask = ImageChops.difference(image, Image.new('RGB', image.size, 'white'))
        box = mask.convert('L').point(lambda v: 255 if v > 25 else 0).getbbox()
        if box:
            a, b, c, d = box
            image = image.crop((max(0, a-10), max(0, b-10), min(image.width, c+10), min(image.height, d+10)))
        scale = min((w - 24) / image.width, (h - 65) / image.height)
        image = image.resize((round(image.width*scale), round(image.height*scale)), Image.Resampling.LANCZOS)
        canvas.paste(image, (x + (w-image.width)//2, y + 60 + (h-65-image.height)//2))
    buf = io.BytesIO()
    canvas.save(buf, format='PNG', dpi=(300, 300))
    return buf.getvalue()


def add_validation_page(doc, title, png, caption, first=False):
    from docx.shared import Mm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    section = doc.sections[-1]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(10)
    section.left_margin = section.right_margin = Mm(10)
    if not first:
        doc.add_page_break()
    p = doc.add_heading(title, level=3)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(io.BytesIO(png), width=Mm(180))
    picture._inline.docPr.set('descr', 'Five calculated panels: protein ribbon, solvent exposure, '
                             'residue counts, Ramachandran plot, polar and apolar accessible area')
    p = doc.add_paragraph(caption)
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        run.font.size = Pt(9)


def build_panels(pdb_text, rama, browser, sasa=None):
    """Render actual stored coordinates and computed values, never sample data."""
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt
    import ramachandran
    from report_writer import svg_to_png, _png_pose_view
    sasa = sasa if sasa is not None else solvent_exposure(pdb_text)
    panels, errors = {}, []

    def chart(key, width, height, render):
        fig, ax = plt.subplots(figsize=(width, height), dpi=180)
        try:
            render(ax)
            ax.spines[['top', 'right']].set_visible(False)
            fig.tight_layout()
            out = io.BytesIO()
            fig.savefig(out, format='png', facecolor='white')
            panels[key] = out.getvalue()
        finally:
            plt.close(fig)

    counts = rama.get('counts', {})
    def count_plot(ax):
        labels = ['Favoured', 'Allowed', 'Outlier', 'Unscored']
        values = [counts.get(k, 0) for k in ('favoured', 'allowed', 'outlier')] + [len(rama.get('unscored', []))]
        bars = ax.bar(labels, values, color=['#ef4444', '#facc15', '#9333ea', '#94a3b8'])
        ax.bar_label(bars, padding=3)
        ax.set_ylabel('Residues'); ax.tick_params(axis='x', labelrotation=25)
        ax.margins(y=.2)
    if not rama.get('_error'):
        chart('count', 5, 3.5, count_plot)
    else:
        errors.append(rama['_error'])
    if not rama.get('_error') and browser:
        panels['rama'] = svg_to_png(ramachandran.plot_svg(rama, width=1000), browser,
                                   width=1200, height=1200, pad=8)
    if pdb_text and browser:
        try:
            with tempfile.TemporaryDirectory(prefix='mumo-ribbon-') as tmp:
                pdb = Path(tmp) / 'protein.pdb'
                pdb.write_text(pdb_text, encoding='utf-8')
                panels['protein'] = _png_pose_view(str(pdb), {}, browser, options={
                    'protein_color': 'spectrum', 'protein_opacity': 1, 'view_mode': 'full',
                    'show_interactions': False, 'show_residues': False, 'show_labels': False,
                    'show_controls': False, 'show_legend': False}, width=700, height=470)
        except Exception as exc:
            errors.append(f'Protein ribbon: {exc}')
    if sasa.get('status') == 'calculated':
        residues = sasa['residues']
        x = list(range(len(residues)))
        ticks = sorted(set([0, len(x)//2, len(x)-1]))
        def axis(ax):
            ax.set_xticks(ticks, [f"{residues[i]['chain']}:{residues[i]['residue']}" for i in ticks])
            ax.set_xlabel('Chain and residue'); ax.set_ylabel('SASA (Å²)')
        def area(ax):
            ax.plot(x, [r['total'] for r in residues], color='#16a34a', linewidth=1.5)
            axis(ax)
        def polar(ax):
            ax.stackplot(x, [r['polar'] for r in residues], [r['apolar'] for r in residues],
                         colors=['#0284c7', '#f59e0b'], labels=['Polar', 'Apolar'])
            axis(ax); ax.legend(loc='upper right', frameon=False)
        chart('exposure', 5, 3.5, area)
        chart('polar', 12, 3.1, polar)
    else:
        errors.append(sasa.get('_error', 'Solvent exposure unavailable'))
    return compose_panels(panels), sasa, errors
