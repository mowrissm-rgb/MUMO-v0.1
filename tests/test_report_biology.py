import io
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))


class ReportBiologyTests(unittest.TestCase):
    def test_mixed_run_report_identifies_old_pairs_and_missing_contacts(self):
        import pandas as pd
        from docx import Document
        import report_writer
        sample = {'rdf': pd.DataFrame([{'Ligand': 'ethanol', 'Target': 'EGFR',
            'SMILES': 'CCO', 'Best affinity (kcal/mol)': -6, 'H-bonds': None,
            'Calculation origin': 'reused', 'Original run ID': 'old-run-123',
            'Original calculation date': '2026-09-30T12:00:00+00:00'}]),
            'meta': {'gene': 'EGFR', 'cache': {'status': 'reuse', 'reused_pairs': 1}}}
        with patch.object(report_writer, 'new_browser', side_effect=RuntimeError('offline')), \
             patch('agents.qm_analyst.xtb_available', return_value=False):
            doc = Document(io.BytesIO(report_writer.build_docking_docx(sample)))
        text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(
            c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn('old-run-123', text)
        self.assertIn('2026-09-30', text)
        self.assertIn('Previously calculated', text)
        self.assertTrue(any(c.text == 'Not calculated' for t in doc.tables for r in t.rows for c in r.cells))

    def test_network_snapshot_embeds_as_an_image(self):
        import pandas as pd
        from docx import Document
        from PIL import Image
        import report_writer
        png = io.BytesIO()
        Image.new('RGB', (200, 100), 'green').save(png, format='PNG')
        class Browser:
            def close(self): pass
        class Manager:
            def stop(self): pass
        sample = {'rdf': pd.DataFrame([{'Ligand': 'ethanol', 'Target': 'EGFR',
                        'SMILES': 'CCO', 'Best affinity (kcal/mol)': 'FAILED'}]),
                  'meta': {'gene': 'EGFR', 'biology': {'string': {'groups': [
                    {'species': 9606, 'result': {'network_svg': '<svg></svg>'}}]}}}}
        with patch.object(report_writer, 'new_browser', return_value=(Manager(), Browser())), \
             patch.object(report_writer, 'svg_to_png', return_value=png.getvalue()), \
             patch('agents.qm_analyst.xtb_available', return_value=False):
            doc = Document(io.BytesIO(report_writer.build_docking_docx(sample)))
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertTrue(any('STRING functional association network' in p.text for p in doc.paragraphs))

    def test_stored_biology_follows_order_and_never_recalculates_on_download(self):
        import pandas as pd
        from docx import Document
        import report_writer
        sample = {'rdf': pd.DataFrame([{'Ligand': 'ethanol', 'Target': 'EGFR',
                    'SMILES': 'CCO', 'Best affinity (kcal/mol)': 'FAILED'}]),
                  'meta': {'gene': 'EGFR', 'docking_mode': 'High', 'exhaustiveness': 32,
                           'replicas': 3, 'biology': {
                    'string': {'groups': [{'species': 9606, 'status': 'available',
                        'mapping': {'EGFR': 'P00533'}, 'created_at': '2026-10-05',
                        'enrichment_scope': 'Input proteins and functional partners',
                        'result': {'partners': [], 'enrichment': [], 'network_svg': ''}}]},
                    'adme': {'compounds': [{'labels': ['ethanol'], 'targets': ['EGFR'],
                        'smiles': 'CCO', 'descriptors': {'MW (g/mol)': 46.069},
                        'ml': {'_error': 'model unavailable'}}], 'no_candidate': []}}}}
        with patch.object(report_writer, 'new_browser', side_effect=RuntimeError('offline')), \
             patch('agents.qm_analyst.xtb_available', return_value=False), \
             patch('docking_biology.collect_string', side_effect=AssertionError('Must not query STRING')), \
             patch('docking_biology.analyze_lead_adme', side_effect=AssertionError('Must not rerun ADME')):
            for _ in range(2):
                doc = Document(io.BytesIO(report_writer.build_docking_docx(sample)))
                headings = [p.text for p in doc.paragraphs if p.style.name == 'Heading 2']
                wanted = ['Ramachandran and target validation', 'Target STRING context',
                          'Ligand HOMO LUMO', 'Docking results', 'Top docking candidate ADME']
                self.assertEqual([h for h in headings if h in wanted], wanted)
                text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(
                    c.text for t in doc.tables for r in t.rows for c in r.cells)
                self.assertIn('P00533', text)
                self.assertIn('46.069', text)
                self.assertIn('High', text)


if __name__ == '__main__':
    unittest.main()
