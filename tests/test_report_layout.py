"""Catch report ordering, portrait sizing and panel omissions in actual DOCX output."""
import io
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))


class ReportLayoutTests(unittest.TestCase):
    def test_report_keeps_methods_before_ordered_results(self):
        import pandas as pd
        from docx import Document
        import report_writer
        data = {'rdf': pd.DataFrame([{'Ligand': 'Aspirin',
                'Best affinity (kcal/mol)': 'FAILED', 'SMILES': 'CC(=O)Oc1ccccc1C(O)=O'}]),
                'meta': {'gene': 'TEST', 'exhaustiveness': 16, 'replicas': 3}}
        # Browser / xTB are external engines; all document assembly stays real.
        with patch.object(report_writer, 'new_browser', side_effect=RuntimeError('test offline')), \
             patch('agents.qm_analyst.xtb_available', return_value=False):
            doc = Document(io.BytesIO(report_writer.build_docking_docx(data)))
        heads = [p.text for p in doc.paragraphs if p.style.name == 'Heading 1']
        expected = ['Methodology', 'Results', 'Discussion', 'Conclusion', 'References']
        self.assertEqual([x for x in heads if x in expected], expected)
        subheads = [p.text for p in doc.paragraphs if p.style.name == 'Heading 2']
        for heading in ['Ramachandran and target validation', 'Ligand HOMO LUMO', 'Docking results']:
            self.assertIn(heading, subheads)
        indexes = [subheads.index(x) for x in ['Ramachandran and target validation',
                   'Ligand HOMO LUMO', 'Docking results']]
        self.assertEqual(indexes, sorted(indexes))
        for section in doc.sections:
            self.assertAlmostEqual(section.page_width / 36000, 210, delta=0.1)
            self.assertAlmostEqual(section.page_height / 36000, 297, delta=0.1)
        self.assertNotIn('homology modeling', '\n'.join(p.text for p in doc.paragraphs).lower())

    def test_five_panels_are_embedded_in_portrait_page(self):
        from docx import Document
        from PIL import Image
        import importlib.util
        self.assertIsNotNone(importlib.util.find_spec('target_validation'),
                             'Five-panel report assembly is missing')
        import target_validation
        panels = {}
        colors = ['#ff0000', '#00aa00', '#0000ff', '#ffaa00', '#aa00ff']
        for name, color in zip(['protein', 'exposure', 'count', 'rama', 'polar'], colors):
            buf = io.BytesIO()
            Image.new('RGB', (400, 300), color).save(buf, format='PNG')
            panels[name] = buf.getvalue()
        png = target_validation.compose_panels(panels)
        im = Image.open(io.BytesIO(png))
        self.assertGreater(im.height, im.width)
        pixels = set(im.convert('RGB').getdata())
        for color in [(255, 0, 0), (0, 170, 0), (0, 0, 255), (255, 170, 0), (170, 0, 255)]:
            self.assertIn(color, pixels)
        doc = Document()
        target_validation.add_validation_page(doc, 'Estrogen receptor alpha [PDB 3ERT]', png,
                                               '245 scored residues', first=True)
        self.assertEqual(len(doc.inline_shapes), 1)
        shape = doc.inline_shapes[0]
        self.assertLess(shape.width, doc.sections[0].page_width)
        self.assertLess(shape.height, doc.sections[0].page_height)
        self.assertAlmostEqual(doc.sections[0].page_width / 36000, 210, delta=.1)

    def test_one_orbital_panel_does_not_reserve_five_empty_cells(self):
        from PIL import Image
        import orbital_figures
        buf = io.BytesIO()
        Image.new('RGB', (620, 1270), 'blue').save(buf, format='PNG')
        try:
            sheets = orbital_figures.contact_sheets(
                [{'label': 'Aspirin', 'surface_png': buf.getvalue()}], compact=True)
        except TypeError:
            self.fail('Compact report sheets are not supported')
        image = Image.open(io.BytesIO(sheets[0]['png'])).convert('RGB')
        # A full-page six-cell canvas would leave 5/6 of this image empty.
        self.assertEqual(image.getbbox(), (0, 0, 620, 1270))
        self.assertEqual(image.getpixel((619, 1269)), (0, 0, 255))


if __name__ == '__main__':
    unittest.main()
