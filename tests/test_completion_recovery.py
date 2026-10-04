import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import subprocess

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))


class CompletionRecoveryTests(unittest.TestCase):
    def test_database_not_found_is_not_retried_or_replaced(self):
        import requests
        import public_database
        response = requests.Response()
        response.status_code = 404
        with patch('requests.get', side_effect=[response, AssertionError('unexpected retry')]):
            self.assertEqual(public_database.get('https://example.org/exact-id').status_code, 404)

    def test_database_certificate_error_is_never_bypassed(self):
        import requests
        import public_database
        with patch('requests.get', side_effect=requests.exceptions.SSLError('certificate invalid')):
            with self.assertRaises(requests.exceptions.SSLError):
                public_database.get('https://example.org/exact-id')

    def test_uniprot_retries_temporary_timeout_without_changing_target(self):
        import requests
        from agents.target_analyst import gene_to_uniprot
        class Response:
            status_code = 200
            def raise_for_status(self): pass
            def json(self): return {'results': [{'primaryAccession': 'P11802'}]}
        with patch('requests.get', side_effect=[requests.Timeout('temporary'), Response()]), \
                patch('time.sleep'):
            self.assertEqual(gene_to_uniprot('CDK4'), 'P11802')

    def test_missing_interaction_count_is_not_printed_as_nan(self):
        import pandas as pd
        from docx import Document
        from report_writer import _add_df_table
        doc = Document()
        _add_df_table(doc, pd.DataFrame([{'Ligand': 'oleic acid', 'H-bonds': float('nan')},
                                       {'Ligand': 'aspirin', 'H-bonds': 0}]))
        self.assertEqual(doc.tables[0].cell(1, 1).text, 'Not calculated')
        self.assertEqual(doc.tables[0].cell(2, 1).text, '0.0')

    def test_failed_target_still_accounts_for_every_requested_pair(self):
        import pipeline_core
        def dock(target, *args, **kwargs):
            if target['gene'] == 'HER2':
                raise RuntimeError('target preparation failed')
            return [{'Ligand': label, 'SMILES': smiles,
                'Best affinity (kcal/mol)': -5.1, 'H-bond residues': '-',
                'All interacting residues': '-', 'H-bonds': None,
                'Total interactions': None, 'Hydrophobic': None}
                for label, smiles in [('aspirin', 'CC(=O)Oc1ccccc1C(=O)O'), ('ethanol', 'CCO')]], {}, {'gene': 'EGFR'}
        with tempfile.TemporaryDirectory() as tmp, patch('pipeline.dock_pipeline', side_effect=dock):
            result = pipeline_core.run_job({'target': ['EGFR', 'HER2'],
                'ligand_objs': [{'label': 'aspirin', 'smiles': 'CC(=O)Oc1ccccc1C(=O)O'},
                                {'label': 'ethanol', 'smiles': 'CCO'}]}, 'vina', tmp, tmp)
        self.assertEqual({(r['Target'], r['Ligand']) for r in result['rows']},
                         {('EGFR', 'aspirin'), ('EGFR', 'ethanol'), ('HER2', 'aspirin'), ('HER2', 'ethanol')})
        self.assertEqual(sum(r['Best affinity (kcal/mol)'] == 'FAILED' for r in result['rows']), 2)

    def test_worker_error_survives_nonzero_exit(self):
        import target_validation
        worker = subprocess.CompletedProcess([], 1,
            json.dumps({'status': 'failed', '_error': 'No protein ATOM coordinates found'}), '')
        with patch.dict('os.environ', {'MUMO_VALIDATION_PYTHON': sys.executable}), \
                patch('target_validation.subprocess.run', return_value=worker):
            result = target_validation.solvent_exposure('')
        self.assertIn('No protein ATOM coordinates', result['_error'])

    def test_flexible_ligand_converges_instead_of_stopping_at_200(self):
        from agents.qm_analyst import prepare_xyz
        xyz, count = prepare_xyz('CCCCCCCC/C=C\\CCCCCCCC(=O)O')
        self.assertGreater(count, 40)
        self.assertEqual(int(xyz.splitlines()[0]), count)

    def test_every_target_retains_own_validation_coordinates(self):
        import pipeline_core
        def dock(target, *args, **kwargs):
            row = {'Ligand': 'aspirin', 'SMILES': 'CC(=O)Oc1ccccc1C(=O)O',
                   'Best affinity (kcal/mol)': -5.1, 'H-bond residues': '-',
                   'All interacting residues': '-', 'H-bonds': None,
                   'Total interactions': None, 'Hydrophobic': None}
            return [row], {}, {'gene': target['gene'], 'validation_pdb_text': target['gene']+' coordinates',
                'center': [1, 2, 3], 'size': [20, 20, 20], 'seed': 42, 'accession': target['gene']+' ID'}
        with tempfile.TemporaryDirectory() as tmp, patch('pipeline.dock_pipeline', side_effect=dock):
            result = pipeline_core.run_job({'target': ['EGFR', 'HER2'],
                'ligand_objs': [{'label': 'aspirin', 'smiles': 'CC(=O)Oc1ccccc1C(=O)O'}]},
                'vina', tmp, tmp)
        for target in ('EGFR', 'HER2'):
            meta = result['meta']['per_target'][target]
            self.assertEqual(meta.get('validation_pdb_text'), target+' coordinates')
            self.assertEqual(meta.get('center'), [1, 2, 3])
            self.assertEqual(meta.get('accession'), target+' ID')
